#!/usr/bin/env python3
"""Produce actual PDF/PPTX/CSV/JSON outputs from a synthetic, unapproved AIFIRST dossier.
Requires reportlab and python-pptx; no network calls, no account access, no real signatures.
"""
import csv, json, sys, pathlib, re
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT=pathlib.Path(sys.argv[2]).resolve()
DATA=json.load(open(sys.argv[1],encoding='utf-8'))
if DATA.get('release_approved') is not False or DATA.get('external_signatures_verified') is not False:
    raise SystemExit('Only unapproved draft dossier allowed')
OUT.mkdir(parents=True,exist_ok=True)
LABEL='SIMULACION / BORRADOR - NO PUBLICAR - SIN APROBACION EXTERNA'
NAVY=colors.HexColor('#1c2f48')
CORAL=colors.HexColor('#ef6754')
MUTED=colors.HexColor('#64748b')
styles={
 'title':ParagraphStyle('title',fontName='Helvetica-Bold',fontSize=24,leading=27,textColor=NAVY,spaceAfter=16),
 'subtitle':ParagraphStyle('subtitle',fontName='Helvetica-Bold',fontSize=13,leading=17,textColor=NAVY,spaceBefore=13,spaceAfter=6),
 'body':ParagraphStyle('body',fontName='Helvetica',fontSize=10,leading=16,textColor=NAVY,spaceAfter=7),
 'warning':ParagraphStyle('warning',fontName='Helvetica-Bold',fontSize=9,leading=14,textColor=CORAL,spaceAfter=15),
}
def P(s,kind='body'):
    return Paragraph(escape(str(s))[:3000],styles[kind])
def make_pdf(fname,title,blocks):
    doc=SimpleDocTemplate(str(OUT/fname),pagesize=A4,topMargin=60,bottomMargin=56,leftMargin=48,rightMargin=48,
                          author='AIFIRST / synthetic-only',title=title)
    elems=[P(title,'title'),P(LABEL,'warning')]
    for heading,lines in blocks:
        elems.append(P(heading,'subtitle'))
        for ln in lines:elems.append(P(ln))
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8);canvas.setFillColor(MUTED)
        canvas.drawString(48,32,'AIFIRST | Draft | Source-linked review required')
        canvas.drawRightString(548,32,str(doc.page))
    doc.build(elems,onFirstPage=footer,onLaterPages=footer)
def output(id):return next((x.get('payload',{}) for x in DATA.get('outputs',[]) if x['id']==id),{})
def limited(items,limit=20):return [str(x)[:450] for x in items][:limit]
ed=DATA.get('edition',{}); title=ed.get('title','Synthetic edition');sub=output('EX03').get('submissions',[])
scenarios=output('EX06').get('scenarios',[])
finance=DATA.get('finance',{});gateList=DATA.get('gates',[])
make_pdf('EX01_CEO_Report_DRAFT.pdf','Executive Discovery Report',[
 ('Context', ['Edition: '+title,'This report is a synthetic draft. No pilot metrics or employee capabilities have been externally verified.']),
 ('Evidence summary',[f'Solutions received: {len(sub)}',f'Scenarios: {len(scenarios)}',f'Recorded external approvals: 0']),
 ('Observed project proposals',[str(s.get('title')) for s in sub] or ['None']),
 ('Risks and decisions',['All gates remain unverified for enterprise release.','External legal, security, financial and jury review required.'])
])
make_pdf('EX06_Scenario_Pack_DRAFT.pdf','Organizational Scenarios',[('Hypothetical alternatives',[
 f"{s.get('name','Untitled')}: {s.get('cost_low','?')}–{s.get('cost_high','?')} EUR; risk: {s.get('risk','Unassessed')}" for s in scenarios] or ['Three scenarios required for review']),
 ('Decision control',['This is not an automated workforce assignment. G5 human approval remains pending.'])])
make_pdf('EX08_Technical_Risk_DRAFT.pdf','Technical and Risk Annex',[
 ('Quarantined artifacts',[f"{a.get('name','artifact')}: {a.get('scan','unknown')}; {a.get('sha','')}" for a in output('EX08').get('quarantined',[])] or ['No artifacts']),
 ('Assurance',['An AV state or local jury flag is not an enterprise signature.','Document confidentiality, SSO, DPA and impact verification still needed.'])
])

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
for h,points in [
 ('AIFIRST / Executive Discovery', ['SYNTHETIC DRAFT  / NOT CLIENT APPROVED',title,'No employee rankings or verified ROI']),
 ('Solution portfolio', [f"{len(sub)} synthetic submissions"]+limited([x.get('title') for x in sub],5)),
 ('Observed capabilities',['Aggregated only; no individual scores','Participation is voluntary and not representative']),
 ('Scenarios',[f"{len(scenarios)} candidate scenarios"]+limited([s.get('name') for s in scenarios],5)),
 ('Activation',['D30: baseline + security','D60: optional controlled pilot','D90: review and decision']),
 ('Economics',['Cost estimate is not an audited expense',f"Inference estimate EUR {finance.get('compute_cost_estimate_eur',0)}"]),
 ('Assurance & gates',['G1 + G3–G7 require verified signatures','No commercial claims without source and human review'])]:
    slide=prs.slides.add_slide(prs.slide_layouts[6]);
    bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor(247,249,252)
    box=slide.shapes.add_textbox(Inches(.8),Inches(.6),Inches(11.6),Inches(1.0));tf=box.text_frame
    tf.text=h;tf.paragraphs[0].runs[0].font.size=Pt(31);tf.paragraphs[0].runs[0].font.bold=True;tf.paragraphs[0].runs[0].font.color.rgb=RGBColor(27,47,72)
    content=slide.shapes.add_textbox(Inches(1.0),Inches(2),Inches(11),Inches(4.1)).text_frame
    for i,line in enumerate(points):
        para=content.paragraphs[0] if i==0 else content.add_paragraph();para.text='• '+str(line)[:130]
        para.font.size=Pt(19);para.space_after=Pt(17)
    tag=slide.shapes.add_textbox(Inches(.9),Inches(7.1),Inches(11),Inches(.2)).text_frame
    tag.text=LABEL;tag.paragraphs[0].runs[0].font.size=Pt(9)
prs.save(OUT/'EX02_Board_Presentation_DRAFT.pptx')

with open(OUT/'EX03_Solution_Portfolio_DRAFT.csv','w',newline='',encoding='utf-8') as fp:
    w=csv.writer(fp);w.writerow(['id','title','status','unverified']);
    for r in sub:w.writerow([r.get('id'),r.get('title'),r.get('state'),'true'])
with open(OUT/'EX04_Capability_Atlas_AGGREGATED_DRAFT.csv','w',newline='',encoding='utf-8') as fp:
    w=csv.writer(fp);w.writerow(['type','suppressed','coverage_note']);w.writerow(['No reportable aggregate','true','Insufficient approved evidence / avoid small-cell disclosure'])
(OUT/'EX05_Knowledge_Graph_DRAFT.json').write_text(json.dumps({**output('EX05'),'state':'PROPOSED_UNVERIFIED'},indent=2,ensure_ascii=False),encoding='utf-8')
with open(OUT/'EX07_Roadmap_30_60_90_DRAFT.csv','w',newline='',encoding='utf-8') as fp:
    w=csv.writer(fp);w.writerow(['milestone','description','gate']);
    for d,a in [(30,'Baseline and legal review'),(60,'Security-approved pilot'),(90,'Sponsor decision')]:w.writerow([f'D{d}',a,'UNVERIFIED'])
(OUT/'EX09_Private_Feedback_NOT_RELEASED.json').write_text(json.dumps({'status':'WITHHELD','reason':'Requires participant authentication, independent assessment and rights review'},indent=2),encoding='utf-8')
(OUT/'EX10_Activation_Bundle_DRAFT.json').write_text(json.dumps({'status':'CANDIDATE','release_allowed':False,'milestones':['D30','D60','D90'],'source':'synthetic'},indent=2),encoding='utf-8')
(OUT/'EX11_Exit_Audit_DRAFT.json').write_text(json.dumps({'status':'NOT_APPROVED','gates':gateList,'retention':'PENDING_CLIENT_POLICY','signed':False},indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'generated':[x.name for x in sorted(OUT.iterdir())],'release_allowed':False,'warning':LABEL}))
